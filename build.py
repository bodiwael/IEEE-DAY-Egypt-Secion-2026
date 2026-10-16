#!/usr/bin/env python3
"""Reads Master_Sheet.xlsx -> builds site/ (encrypted attendee data), backend/Code.gs (holds the data key + staff code hashes), staff_codes.csv
   Usage: pip install openpyxl cryptography   then   python3 build.py [Master_Sheet.xlsx]
   Optional: put your Apps Script /exec URL in api_url.txt so phones need no ?api= link.
   ALWAYS deploy site/ AND Code.gs together - each build has a new data key."""
import openpyxl,json,os,re,base64,hashlib,secrets,csv,sys,shutil,time
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
H=os.path.dirname(os.path.abspath(__file__));XL=sys.argv[1] if len(sys.argv)>1 else H+'/Master_Sheet.xlsx'
SALT='ieee26fy';CATS={'VIP':'VIP','Former Officers':'FO','YP':'YP','SAC':'SAC','Ambassador':'AMB','Attendees':'AT','Volunteers':'VOL'}
ROLES={'bus','reg','kit','joker','admin'}
A={'id':['id','ticket id'],'name':['name'],'email':['email'],'point':['point'],'branch':['student branch','branch'],'nid':['national id','nationalid','nid'],'tshirt':['t-shirt size','tshirt','t-shirt','size'],'phone':['phone','mobile']}
b=lambda x:base64.b64encode(x).decode();nz=lambda s:re.sub(r'[^A-Z0-9]','',str(s).upper())
def s(v):
    if v is None:return ''
    if isinstance(v,float) and v.is_integer():v=int(v)
    return str(v).strip()
wb=openpyxl.load_workbook(XL,data_only=True);P=[];warn=[];seen={}
for ws in wb:
    if ws.title in('Staff','READ ME'):continue
    rows=list(ws.iter_rows(values_only=True));hd=[s(x).lower() for x in rows[0]]
    ix={k:next((hd.index(a) for a in v if a in hd),None) for k,v in A.items()}
    miss=[k for k in('id','name','email') if ix[k] is None]
    if miss:warn.append(f'sheet {ws.title}: missing column {miss}');continue
    for k in('nid','tshirt','phone'):
        if ix[k] is None:warn.append(f'sheet {ws.title}: no {k} column - header must be one of {A[k]}')
    g=lambda r,k:s(r[ix[k]]) if ix[k] is not None and ix[k]<len(r) else ''
    cat=CATS.get(ws.title,ws.title[:3].upper())
    for r in rows[1:]:
        if not g(r,'id') and not g(r,'name'):continue
        i=g(r,'id').zfill(3);u=cat+i;pt=g(r,'point');pt=pt if pt in('Ramsis','October') else 'No'
        nid=re.sub(r'[\s\u200b-\u206f\ufeff]','',g(r,'nid'))
        if re.search(r'E\+',nid):warn.append(f'{ws.title} {i}: National ID damaged by Excel (scientific notation) - format column as Text')
        elif not nid:warn.append(f'{ws.title} {i} {g(r,"name")}: National ID missing')
        elif not re.fullmatch(r'\d{14}',nid):warn.append(f'{ws.title} {i} {g(r,"name")}: National ID "{nid}" is not 14 digits')
        if not g(r,'tshirt'):warn.append(f'{ws.title} {i} {g(r,"name")}: T-shirt size missing')
        em=g(r,'email').lower()
        if not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$',em):warn.append(f'{ws.title} {i}: invalid email "{em}"')
        if u in seen:warn.append(f'DUPLICATE ticket {u}')
        seen[u]=1;P.append([u,i,' '.join(g(r,'name').split()),em,pt,cat,g(r,'branch'),nid,g(r,'tshirt').upper(),g(r,'phone')])
sheet=wb['Staff'];T={};srows=list(sheet.iter_rows(values_only=True));sh=[s(x).lower() for x in srows[0]]
cv=lambda r,n:s(r[sh.index(n)]) if n in sh and sh.index(n)<len(r) else ''
for r in srows[1:]:
    if not s(r[0]):continue
    sid=nz(r[0]);role=cv(r,'role').lower()
    if role not in ROLES:warn.append(f'Staff {sid}: bad role "{role}" (use {sorted(ROLES)})');continue
    if role=='bus' and not cv(r,'bus'):warn.append(f'Staff {sid}: bus lead without a Bus number');continue
    T[sid]=dict(n=cv(r,'name'),r=role,bus=cv(r,'bus'),pt=cv(r,'point'),e=cv(r,'email'),ph=cv(r,'phone'))
codes={}
if os.path.exists(H+'/staff_codes.csv'):codes={r['staff_id']:r['code'] for r in csv.DictReader(open(H+'/staff_codes.csv')) if re.fullmatch(r'\d{6}',r['code'])}
for sid in T:codes.setdefault(sid,''.join(secrets.choice('0123456789') for _ in range(6)))
K=os.urandom(32);iv=os.urandom(12);BUILD=time.strftime('%d%b-%H%M')
DATA=[b(iv),b(AESGCM(K).encrypt(iv,json.dumps({'P':P,'T':T},separators=(',',':'),ensure_ascii=False).encode(),None))]
API=open(H+'/api_url.txt').read().strip() if os.path.exists(H+'/api_url.txt') else ''
os.makedirs(H+'/site',exist_ok=True)
t=open(H+'/template.html').read().replace('__DATA__',json.dumps(DATA)).replace('__SALT__',SALT).replace('__BUILD__',BUILD).replace('__API__',API).replace('__BANNER__','data:image/jpeg;base64,'+b(open(H+'/assets/banner.jpg','rb').read()))
open(H+'/site/index.html','w').write(t)
for f in('sw.js','manifest.json','icon-192.png','icon-512.png'):shutil.copy(H+'/assets/'+f,H+'/site/'+f)
open(H+'/site/.nojekyll','w').write('')
hs={sid:hashlib.sha256((SALT+sid+codes[sid]).encode()).hexdigest() for sid in T}
open(H+'/backend/Code.gs','w').write(open(H+'/backend/Code.gs.tpl').read().replace('__STAFF__',json.dumps(hs)).replace('__KEY__',b(K)).replace('__BUILD__',BUILD))
w=csv.writer(open(H+'/staff_codes.csv','w',newline=''));w.writerow(['staff_id','name','role','bus','phone','code'])
for sid,x in T.items():w.writerow([sid,x['n'],x['r'],x['bus'],x['ph'],codes[sid]])
nf=sum(1 for p in P if p[7]);tf=sum(1 for p in P if p[8]);pf=sum(1 for p in P if p[9])
print(f'OK build {BUILD}: {len(P)} people, {len(T)} staff. NID filled:{nf} T-shirt filled:{tf} phone:{pf}')
if P and (nf==0 or tf==0):print('!!! ERROR: National ID / T-shirt columns are EMPTY - check the sheet headers and that you saved the file. The app will show "—".')
for x in warn[:40]:print('WARNING:',x)
print(f'({len(warn)} warnings)')
print('NEXT: paste backend/Code.gs into Apps Script > Deploy > Manage deployments > edit > New version, AND re-upload site/. Both are needed.' + ('' if API else '  (api_url.txt not set: open the site once with ?api=YOUR_EXEC_URL)'))
